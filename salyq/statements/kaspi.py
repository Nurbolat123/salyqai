"""Парсер выписки по счёту Kaspi Business (PDF, XLSX или CSV).

Kaspi периодически меняет вид выгрузки, поэтому колонки ищутся по набору
синонимов, а строка заголовка — автоматически. Поддерживаются оба варианта
сумм: раздельные «Дебет/Кредит» и одна колонка «Сумма» со знаком.

Парсер не знает ключей шифрования: он выдаёт dedup_key (открытый текст, живёт
только в памяти), а репозиторий превращает его в HMAC-отпечаток.
"""

import csv
import io
import re
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from typing import Any, Iterable, Literal

from salyq.money import MoneyError, to_tiyn

BANK = "kaspi"
Direction = Literal["in", "out"]


class StatementParseError(ValueError):
    pass


COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("дата операции", "дата", "дата проводки", "дата и время", "күні", "date"),
    "reference": (
        "номер документа", "№ документа", "номер док", "№ док", "референс", "reference",
        "номер операции", "құжат нөмірі",
    ),
    "debit": ("дебет", "списание", "расход", "сумма списания", "дебет (kzt)", "дебет, kzt"),
    "credit": ("кредит", "поступление", "приход", "сумма поступления", "кредит (kzt)", "кредит, kzt"),
    "amount": ("сумма", "сумма операции", "сумма (kzt)", "сумма, kzt", "сома"),
    "currency": ("валюта", "валюта операции", "currency"),
    "counterparty": (
        "контрагент", "наименование контрагента", "получатель/отправитель",
        "отправитель/получатель", "наименование получателя/отправителя",
        "наименование бенефициара/отправителя", "корреспондент",
    ),
    "counterparty_id": ("иин/бин", "иин/бин контрагента", "бин/иин", "бин/иин контрагента", "жсн/бсн"),
    "counterparty_account": ("счет контрагента", "iban контрагента", "иик контрагента", "иик"),
    "knp": ("кнп", "код назначения платежа"),
    "purpose": ("назначение платежа", "назначение", "детали", "описание", "детали операции", "төлем мақсаты"),
    "operation": ("операция", "тип операции", "вид операции"),
}

_REQUIRED_ANY_AMOUNT = (("debit", "credit"), ("amount",))
_IBAN_RE = re.compile(r"(?<![A-Z0-9])KZ\d{2}(?:\s?[A-Z0-9]){16}(?![A-Z0-9])")
_DATE_RE = re.compile(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})")
_ISO_DATE_RE = re.compile(r"(\d{4})-(\d{2})-(\d{2})")
_PERIOD_RE = re.compile(
    r"период\D{0,20}?(\d{1,2}[./-]\d{1,2}[./-]\d{2,4})\s*(?:-|–|—|по)\s*(\d{1,2}[./-]\d{1,2}[./-]\d{2,4})",
    re.IGNORECASE,
)
_SUMMARY_PREFIXES = ("итого", "всего", "остаток", "обороты", "барлығы")


@dataclass(frozen=True)
class ParsedTransaction:
    op_date: date
    amount_tiyn: int  # всегда > 0, направление — в direction
    direction: Direction
    currency: str
    reference: str
    counterparty: str
    counterparty_id: str
    counterparty_account: str
    knp: str
    purpose: str
    operation: str
    row_number: int  # номер строки в исходном файле (1-based)
    dedup_key: str = ""


@dataclass
class ParsedStatement:
    bank: str
    account_iban: str | None
    period_from: date | None
    period_to: date | None
    transactions: list[ParsedTransaction]
    skipped_rows: list[tuple[int, str]] = field(default_factory=list)  # (строка, причина)
    duplicates_in_file: int = 0


def _norm_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower().replace("ё", "е")


def _norm_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _match_columns(header: list[Any]) -> dict[str, int] | None:
    normalized = [_norm_header(h) for h in header]
    mapping: dict[str, int] = {}
    for field_name, aliases in COLUMN_ALIASES.items():
        wanted = {a.replace("ё", "е") for a in aliases}
        for idx, h in enumerate(normalized):
            if h in wanted and idx not in mapping.values():
                mapping[field_name] = idx
                break
    if "date" not in mapping:
        return None
    if not any(all(f in mapping for f in group) for group in _REQUIRED_ANY_AMOUNT):
        return None
    return mapping


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = _norm_text(value)
    if m := _DATE_RE.search(s):
        d, mth, y = (int(x) for x in m.groups())
        if y < 100:
            y += 2000
        return date(y, mth, d)
    if m := _ISO_DATE_RE.search(s):
        y, mth, d = (int(x) for x in m.groups())
        return date(y, mth, d)
    raise StatementParseError(f"не удалось разобрать дату: {value!r}")


def _parse_amount(value: Any) -> int:
    """Сумма ячейки → минимальные единицы со знаком. Пустая ячейка → 0."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0
    if isinstance(value, bool):
        raise StatementParseError(f"некорректная сумма: {value!r}")
    if isinstance(value, int):
        return value * 100
    if isinstance(value, float):
        # Excel хранит числа как float; до 2 знаков repr точен для сумм выписки
        value = repr(round(value, 2))
    s = str(value).strip()
    negative = s.startswith(("-", "−")) or (s.startswith("(") and s.endswith(")"))
    s = s.strip("()").lstrip("-−+").strip()
    # «1.234.567,89» → «1234567,89»
    if "," in s and "." in s:
        s = s.replace(".", "") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    try:
        tiyn = to_tiyn(s)
    except MoneyError as exc:
        raise StatementParseError(str(exc)) from exc
    return -tiyn if negative else tiyn


def _cell(row: list[Any], mapping: dict[str, int], name: str) -> Any:
    idx = mapping.get(name)
    return row[idx] if idx is not None and idx < len(row) else None


def dedup_key(bank: str, account: str | None, tx: ParsedTransaction, occurrence: int) -> str:
    """Ключ дедупликации (ТЗ 4.2: банк, дата, сумма, референс).

    С референсом ключ не зависит от текстовых полей, поэтому одна и та же операция
    из PDF и XLSX выгрузок совпадёт, даже если банк по-разному обрезал назначение.
    Без референса ключ строится по содержимому строки плюс порядковый номер среди
    полностью одинаковых строк файла: две реально разные одинаковые операции не
    схлопываются, а повторная загрузка — схлопывается.
    """
    base = [bank, (account or "").upper(), tx.op_date.isoformat(), tx.direction, tx.currency, str(tx.amount_tiyn)]
    if tx.reference:
        parts = base + ["ref", tx.reference.lower()]
    else:
        parts = base + [
            "row", tx.counterparty_id, tx.counterparty_account.upper().replace(" ", ""),
            tx.counterparty.lower(), tx.knp, tx.purpose.lower(), str(occurrence),
        ]
    return "\x1f".join(parts)


def _find_period(text: str) -> tuple[date, date] | None:
    if m := _PERIOD_RE.search(text):
        try:
            return _parse_date(m.group(1)), _parse_date(m.group(2))
        except (StatementParseError, ValueError):
            return None
    return None


def parse_rows(rows: Iterable[list[Any]], bank: str = BANK) -> ParsedStatement:
    rows = list(rows)
    account: str | None = None
    period: tuple[date, date] | None = None
    mapping: dict[str, int] | None = None
    header_idx = -1
    for i, row in enumerate(rows):
        line = " ".join(_norm_text(c) for c in row)
        if account is None and (m := _IBAN_RE.search(line.upper())):
            account = m.group().replace(" ", "")
        period = period or _find_period(line)
        mapping = _match_columns(row)
        if mapping:
            header_idx = i
            break
    if mapping is None:
        raise StatementParseError("не найдена строка заголовка с датой и суммами — это не выписка Kaspi Business?")

    result = ParsedStatement(bank, account, None, None, transactions=[])
    occurrences: dict[tuple, int] = {}
    seen_refs: set[tuple[str, str, int]] = set()

    for i, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(_norm_text(c) for c in row):
            continue
        if _norm_header(row[0] if row else "").startswith(_SUMMARY_PREFIXES):
            continue
        if _match_columns(row):  # заголовок, повторённый на каждой странице PDF
            continue
        try:
            op_date = _parse_date(_cell(row, mapping, "date"))
        except (StatementParseError, ValueError) as exc:
            result.skipped_rows.append((i, str(exc)))
            continue
        try:
            if "amount" in mapping and not ("debit" in mapping and "credit" in mapping):
                signed = _parse_amount(_cell(row, mapping, "amount"))
            else:
                signed = _parse_amount(_cell(row, mapping, "credit")) - abs(
                    _parse_amount(_cell(row, mapping, "debit"))
                )
        except StatementParseError as exc:
            result.skipped_rows.append((i, str(exc)))
            continue
        if signed == 0:
            result.skipped_rows.append((i, "нулевая сумма"))
            continue

        tx = ParsedTransaction(
            op_date=op_date,
            amount_tiyn=abs(signed),
            direction="in" if signed > 0 else "out",
            currency=(_norm_text(_cell(row, mapping, "currency")).upper().replace("₸", "KZT") or "KZT"),
            reference=_norm_text(_cell(row, mapping, "reference")),
            counterparty=_norm_text(_cell(row, mapping, "counterparty")),
            counterparty_id=re.sub(r"\D", "", _norm_text(_cell(row, mapping, "counterparty_id"))),
            counterparty_account=_norm_text(_cell(row, mapping, "counterparty_account")),
            knp=_norm_text(_cell(row, mapping, "knp")),
            purpose=_norm_text(_cell(row, mapping, "purpose")),
            operation=_norm_text(_cell(row, mapping, "operation")),
            row_number=i,
        )

        # Явный дубль внутри файла: тот же референс, дата и сумма
        if tx.reference:
            ref_key = (tx.reference, tx.op_date.isoformat(), signed)
            if ref_key in seen_refs:
                result.duplicates_in_file += 1
                continue
            seen_refs.add(ref_key)

        content = (tx.op_date, tx.direction, tx.currency, tx.amount_tiyn, tx.reference, tx.counterparty_id,
                   tx.counterparty_account, tx.counterparty.lower(), tx.knp, tx.purpose.lower())
        occ = occurrences.get(content, 0)
        occurrences[content] = occ + 1
        result.transactions.append(replace(tx, dedup_key=dedup_key(bank, account, tx, occ)))

    dates = [t.op_date for t in result.transactions]
    if period:
        result.period_from, result.period_to = period
    elif dates:
        result.period_from, result.period_to = min(dates), max(dates)
    return result


def _read_csv(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise StatementParseError("неизвестная кодировка CSV")
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=";,\t").delimiter
    except csv.Error:
        delimiter = ";"
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _read_xlsx(data: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def _read_pdf(data: bytes) -> list[list[Any]]:
    """Шапка (текст над первой таблицей: счёт, период) + строки всех таблиц документа."""
    import pdfplumber

    rows: list[list[Any]] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page_no, page in enumerate(pdf.pages):
                tables = page.find_tables()
                if page_no == 0:
                    top = tables[0].bbox[1] if tables else page.height
                    header = page.crop((0, 0, page.width, top)).extract_text() if top > 0 else ""
                    rows.extend([line] for line in (header or "").splitlines())
                for table in tables:
                    rows.extend([_norm_text(c) for c in r] for r in table.extract())
    except Exception as exc:  # pdfminer бросает разнородные исключения на битых файлах
        raise StatementParseError(f"не удалось прочитать PDF: {exc}") from exc
    return rows


def parse_kaspi_statement(data: bytes, filename: str = "") -> ParsedStatement:
    name = filename.lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        rows = _read_pdf(data)
    elif name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        rows = _read_xlsx(data)
    elif name.endswith((".csv", ".txt")) or not name:
        rows = _read_csv(data)
    else:
        raise StatementParseError(f"формат {filename!r} не поддерживается: загрузите PDF, XLSX или CSV")
    return parse_rows(rows)
